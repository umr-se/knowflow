import csv
import json
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from xlrd import open_workbook
from pptx import Presentation
from pypdf import PdfReader

from ..config import get_settings
from .chunking import chunk_text


SUPPORTED_EXTENSIONS = {
    ".pdf", ".txt", ".md", ".csv", ".tsv", ".json", ".jsonl",
    ".xlsx", ".xls", ".xlsm", ".docx", ".pptx", ".html", ".htm", ".xml",
}


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> list[dict]:
    reader = PdfReader(BytesIO(data))
    items = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            items.append({"text": text, "page_number": page_number})
    return items


def _extract_xlsx(data: bytes) -> list[dict]:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                parts.append(" | ".join(values))
    return [{"text": "\n".join(parts), "page_number": None}]


def _extract_xls(data: bytes) -> list[dict]:
    """Extract legacy Excel .xls workbooks without changing the RAG pipeline."""
    workbook = open_workbook(file_contents=data)
    sections: list[dict] = []

    for sheet in workbook.sheets():
        parts = [f"Sheet: {sheet.name}"]
        for row_index in range(sheet.nrows):
            values = []
            for col_index in range(sheet.ncols):
                value = sheet.cell_value(row_index, col_index)
                if value is None:
                    value = ""
                values.append(str(value).strip())
            if any(value for value in values):
                parts.append(" | ".join(values))

        text = "\n".join(parts).strip()
        if text:
            sections.append({"text": text, "page_number": None})

    return sections


def _extract_delimited(data: bytes, delimiter: str) -> list[dict]:
    text = _decode_text(data)
    reader = csv.reader(StringIO(text), delimiter=delimiter)
    rows = [" | ".join(cell.strip() for cell in row) for row in reader]
    return [{"text": "\n".join(rows), "page_number": None}]


def _extract_docx(data: bytes) -> list[dict]:
    doc = DocxDocument(BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [{"text": "\n".join(parts), "page_number": None}]


def _extract_pptx(data: bytes) -> list[dict]:
    presentation = Presentation(BytesIO(data))
    items = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            items.append({"text": "\n".join(texts), "page_number": slide_number})
    return items


def _extract_markup(data: bytes) -> list[dict]:
    soup = BeautifulSoup(_decode_text(data), "html.parser")
    return [{"text": soup.get_text("\n", strip=True), "page_number": None}]


def _extract_json(data: bytes, json_lines: bool = False) -> list[dict]:
    text = _decode_text(data)
    if json_lines:
        values = [json.loads(line) for line in text.splitlines() if line.strip()]
        pretty = "\n".join(json.dumps(item, ensure_ascii=False) for item in values)
    else:
        pretty = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    return [{"text": pretty, "page_number": None}]


def build_chunks(filename: str, file_type: str, file_bytes: bytes) -> list[dict]:
    """Extract common local document formats and split them into RAG chunks."""
    settings = get_settings()
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type '{extension or 'unknown'}'. Supported: {supported}")

    if extension == ".pdf":
        sections = _extract_pdf(file_bytes)
    elif extension in {".xlsx", ".xlsm"}:
        sections = _extract_xlsx(file_bytes)
    elif extension == ".xls":
        sections = _extract_xls(file_bytes)
    elif extension == ".csv":
        sections = _extract_delimited(file_bytes, ",")
    elif extension == ".tsv":
        sections = _extract_delimited(file_bytes, "\t")
    elif extension == ".docx":
        sections = _extract_docx(file_bytes)
    elif extension == ".pptx":
        sections = _extract_pptx(file_bytes)
    elif extension in {".html", ".htm", ".xml"}:
        sections = _extract_markup(file_bytes)
    elif extension == ".json":
        sections = _extract_json(file_bytes)
    elif extension == ".jsonl":
        sections = _extract_json(file_bytes, json_lines=True)
    else:
        sections = [{"text": _decode_text(file_bytes), "page_number": None}]

    chunks: list[dict] = []
    for section in sections:
        text = section["text"].strip()
        if not text:
            continue
        for content in chunk_text(text, settings.chunk_size, settings.chunk_overlap):
            chunks.append({"content": content, "page_number": section.get("page_number")})
    return chunks
